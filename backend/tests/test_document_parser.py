"""Tests for the resume document parser (PDF / DOCX / text -> clean text).

DOCX and PDF fixtures are generated in-process (python-docx / fpdf2) so real
extraction paths are exercised, not just dispatch.
"""
import io

import pytest

from app.services.document_parser import (
    DOCX_MIME,
    PDF_MIME,
    DocumentParseError,
    UnsupportedDocumentError,
    extract_text,
)

RESUME_LINES = ["Jane Doe", "Software Engineer", "Skills: Python, SQL, Airflow"]


def _make_docx() -> bytes:
    from docx import Document

    document = Document()
    for line in RESUME_LINES:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _make_pdf() -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    for line in RESUME_LINES:
        pdf.cell(0, 10, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def test_plain_text_passthrough():
    text = extract_text(b"Jane Doe\nPython, SQL", content_type="text/plain", filename="r.txt")
    assert "Jane Doe" in text
    assert "Python, SQL" in text


def test_docx_extraction_by_mime():
    text = extract_text(_make_docx(), content_type=DOCX_MIME, filename="resume.docx")
    assert "Jane Doe" in text
    assert "Python" in text


def test_docx_extraction_by_extension_only():
    # No content-type, just a .docx filename — should still route to the DOCX path.
    text = extract_text(_make_docx(), filename="resume.docx")
    assert "Software Engineer" in text


def test_pdf_extraction():
    text = extract_text(_make_pdf(), content_type=PDF_MIME, filename="resume.pdf")
    assert "Jane" in text
    assert "Python" in text


def test_octet_stream_sniffs_pdf_magic_bytes():
    # Generic content-type + .bin name -> fall back to magic-byte sniffing.
    text = extract_text(_make_pdf(), content_type="application/octet-stream", filename="blob.bin")
    assert "Jane" in text


def test_octet_stream_sniffs_docx_zip_magic():
    text = extract_text(_make_docx(), content_type="application/octet-stream", filename="blob.bin")
    assert "Jane Doe" in text


def test_rtf_extraction():
    rtf = (
        r"{\rtf1\ansi\deff0 Jane Doe\par Software Engineer\par "
        r"Skills: Python, SQL\par}"
    ).encode()
    text = extract_text(rtf, content_type="application/rtf", filename="resume.rtf")
    assert "Jane Doe" in text
    assert "Python" in text


def test_unsupported_type_raises():
    with pytest.raises(UnsupportedDocumentError):
        extract_text(b"\x89PNG\r\n\x1a\n", content_type="image/png", filename="photo.png")


def test_empty_upload_raises():
    with pytest.raises(DocumentParseError):
        extract_text(b"", content_type="text/plain", filename="empty.txt")


def test_oversize_upload_raises():
    big = b"x" * (10 * 1024 * 1024 + 1)
    with pytest.raises(DocumentParseError):
        extract_text(big, content_type="text/plain", filename="big.txt")


def test_whitespace_is_normalized():
    text = extract_text(
        b"Line one\r\n\r\n\r\n\r\nLine    two\t\there",
        content_type="text/plain",
        filename="r.txt",
    )
    assert "\n\n\n" not in text  # 3+ blank lines collapsed
    assert "    " not in text  # runs of spaces collapsed
    assert "Line one" in text and "Line two here" in text


def test_image_only_pdf_raises_parse_error():
    # A structurally-valid-ish PDF header with no extractable text -> clean error,
    # not a 500. (We don't OCR.)
    with pytest.raises(DocumentParseError):
        extract_text(b"%PDF-1.4\n%garbage-no-text", content_type=PDF_MIME, filename="scan.pdf")
